#!/usr/bin/env python3
"""Import supplied examples while preserving original dates and coordinates.

Run from any directory: python /path/to/GridLock/scripts/import_workbook.py
The bundled data/review_annotations.json adds notes only. --review can select a
compatible annotation file; --check verifies current generated JSON without writes.
"""
import argparse
import datetime as dt
import hashlib
import json
import math
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_URL = '/sources/Projects_Overlaps.xlsx'
DESC_PDF_URL = '/sources/Dominion_2024-2028_Project_Descriptions.pdf'
GPC_PDF_URL = '/sources/Georgia_Power_2025_IRP_Volume_3_PUBLIC_DISCLOSURE.pdf'
SHORT_NAMES = {
    'DESC_1': 'Stevens Creek–Hooks',
    'DESC_2': 'Hooks–Thurmond',
    'DESC_3': 'Jasper–Okatie #2',
    'DESC_4': 'Queensboro–Fort Johnson',
    'DESC_5': 'Okatie–Bluffton',
    'GPC_1': 'Evans–Thurmond #5',
    'GPC_2': 'McIntosh reactors',
    'GPC_3': 'Goshen–McIntosh',
    'GPC_4': 'Mitchell–North Tifton',
    'GPC_5': 'Jesup–Ludowici',
}


def normalize_date(value):
    """Return a date-only parsing value without altering the preserved raw value."""
    if value is None:
        return None
    if isinstance(value, str):
        return dt.datetime.strptime(value, '%m/%d/%Y').date().isoformat()
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    raise ValueError(f'Unsupported workbook date type: {type(value).__name__}')


def source_as_of(source, latest):
    """Describe the source's actual date meaning without inventing publication dates."""
    snapshot = latest.get('source_snapshot') or source.get('snapshot_as_of')
    if snapshot:
        return f'Snapshot: {snapshot}'
    report_date = source.get('report_date') or latest.get('report_date')
    if report_date:
        return f'Report dated {report_date}'
    if source.get('document_date'):
        return f'Document dated {source["document_date"]}'
    if source.get('pdf_created'):
        created = source['pdf_created'].split('T')[0]
        return f'Publication date unspecified; PDF creation metadata: {created}'
    return 'Source date not available'


def review_fields(raw, company, note, source_maps, reviewed_on):
    """Keep the research layer separate from the original workbook values."""
    ev = note.get('scope_schedule_evidence', {})
    source = {}
    latest = {}
    source_id = None
    status_source = {}
    status_page = None
    status_date = None
    loc = note.get('coordinate_assessment', {}).get('verdict') or 'Location has not been reviewed.'
    if company == 'DESC':
        latest = ev.get('latest_evidence', {})
        planned = latest.get('latest_exact_planned_in_service')
        actual_month = latest.get('actual_in_service_month')
        source_id = latest.get('schedule_source') if planned else latest.get('reported_in')
        source = source_maps.get('dominion', {}).get(source_id, {})
        status_id = latest.get('reported_in')
        if status_id and status_id != source_id:
            status_source = source_maps.get('dominion', {}).get(status_id, {})
            if not status_source.get('url'):
                raise ValueError(f'{raw["project_id"]}: missing separate status source {status_id}')
            status_page = latest.get('pdf_page')
            status_date = latest.get('evidence_date') or status_source.get('document_date')
        if planned:
            date_note = 'Latest reviewed planned milestone; not confirmed completion.'
        elif actual_month:
            date_note = 'Official completion statement at month precision; do not invent a day.'
        else:
            date_note = 'No reviewed later date is available for this record.'
        review_row = {
            'scope': ev.get('geometry', {}).get('location_detail') or raw['project_name'],
            'locationNote': loc,
            'status': latest.get('status', 'Unreviewed'),
            'latestDate': planned or actual_month,
            'latestDatePrecision': 'day' if planned else 'month' if actual_month else 'unknown',
            'dateNote': date_note,
            'sourceUrl': source.get('url', ''),
            'sourcePage': latest.get('schedule_pdf_page', latest.get('pdf_page', 1)),
            'warnings': [ev[k] for k in ['identity_caution', 'scope_caution'] if ev.get(k)],
        }
        original_source = {
            'title': 'DESC 2024–2028 project descriptions',
            'url': DESC_PDF_URL,
            'page': ev.get('original_page', 1),
            'asOf': '2024–2028 planning horizon',
            'dateNote': 'Publication date unspecified; PDF creation metadata 5 March 2024.',
        }
    elif company == 'GPC':
        latest = ev.get('latest_confirmed_plan', {})
        source_id = latest.get('source')
        source = source_maps.get('georgia_power', {}).get(source_id, {})
        # Page references belong to the bundled PDF, not the official docket landing page.
        source_url = GPC_PDF_URL if source_id == 'irp2025' else source.get('url', '')
        review_row = {
            'scope': ev.get('public_scope', raw['project_name']),
            'locationNote': loc,
            'status': latest.get('status', 'Unreviewed'),
            'latestDate': latest.get('value'),
            'latestDatePrecision': latest.get('precision', 'unknown'),
            'dateNote': ev.get('schedule_change', 'The source field is a need date, not a verified completion or construction interval. Current execution status was not established.'),
            'sourceUrl': source_url,
            'sourcePage': latest.get('pdf_page', ev.get('irp_pdf_page', 1)),
            'warnings': [ev[k] for k in ['coordinate_caution', 'later_search_result'] if ev.get(k)],
        }
        original_source = {
            'title': 'Georgia Power 2025 IRP Volume 3 · Public Disclosure',
            'url': GPC_PDF_URL,
            'page': ev.get('irp_pdf_page', 1),
            'asOf': 'December 2024 snapshot',
            'dateNote': 'The five supplied dates are labeled Need Date in this report, not actual in-service dates.',
        }
    else:
        # New utilities remain unreviewed instead of inheriting Georgia's report or date semantics.
        review_row = {
            'scope': raw['project_name'], 'locationNote': loc, 'status': 'Unreviewed',
            'latestDate': None, 'latestDatePrecision': 'unknown',
            'dateNote': 'Date semantics and utility report have not been reviewed.',
            'sourceUrl': '', 'sourcePage': 1, 'warnings': [],
        }
        original_source = {
            'title': 'Supplied project workbook', 'url': WORKBOOK_URL, 'page': 1,
            'asOf': 'As supplied', 'dateNote': 'The workbook date field has not been reviewed for this utility.',
        }
    if not note and company in ('DESC', 'GPC'):
        original_source['dateNote'] += ' This record’s report page has not been located; the link opens the report at its first page.'
    review_row.update({
        'reviewedOn': reviewed_on if note else None,
        'sourceAsOf': source_as_of(source, latest),
        'statusSourceUrl': status_source.get('url'),
        'statusSourcePage': status_page,
        'statusEvidenceDate': status_date,
    })
    return review_row, original_source


def build_dataset(workbook_path, review):
    """Build the two output objects in memory; the workbook is always read-only."""
    notes = {p['id']: p for p in review.get('projects', [])}
    if len(notes) != len(review.get('projects', [])):
        raise ValueError('Duplicate review annotation IDs')
    source_maps = {
        key: {source['id']: source for source in value.get('sources', [])}
        for key, value in review.get('utility_research', {}).items()
    }
    rows = []
    wb = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        sheet = wb['projects']
        headers = [cell.value for cell in sheet[1]]
        required = {'project_id', 'utility', 'state', 'project_name', 'in_service_date'}
        required.update(f'{field}_{end}' for end in ('a', 'b') for field in ('name', 'lat', 'lon'))
        if not required.issubset(headers):
            raise ValueError(f'Missing project columns: {sorted(required - set(headers))}')
        for cells in sheet.iter_rows(min_row=2, values_only=True):
            raw = dict(zip(headers, cells))
            if not raw['project_id']:
                continue
            pid = str(raw['project_id'])
            company = pid.split('_')[0]
            date = raw['in_service_date']
            endpoints = []
            for key in ('a', 'b'):
                lat, lon = raw['lat_' + key], raw['lon_' + key]
                if (lat is None) != (lon is None):
                    raise ValueError(f'{pid}: incomplete coordinate pair at endpoint {key}')
                if lat is not None:
                    valid_numbers = all(isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v) for v in (lat, lon))
                    if not valid_numbers or not (-90 <= lat <= 90 and -180 <= lon <= 180):
                        raise ValueError(f'{pid}: invalid coordinate at endpoint {key}')
                endpoints.append({'name': raw['name_' + key], 'coordinate': [lat, lon] if lat is not None else None})
            note = notes.get(pid, {})
            review_row, source = review_fields(raw, company, note, source_maps, review.get('reviewed_on'))
            rows.append({
                'id': pid, 'company': company, 'utility': raw['utility'],
                'name': raw['project_name'], 'shortName': SHORT_NAMES.get(pid, raw['project_name'] or pid),
                'state': raw['state'], 'endpoints': endpoints,
                'originalDate': normalize_date(date),
                'originalDateRaw': date.isoformat() if isinstance(date, (dt.datetime, dt.date)) else date,
                'dateMeaning': 'planned_in_service' if company == 'DESC' else 'need_date' if company == 'GPC' else 'unknown',
                'originalSource': source, 'review': review_row,
            })
    finally:
        wb.close()
    if len({row['id'] for row in rows}) != len(rows):
        raise ValueError('Duplicate project IDs')
    metadata = {
        'title': 'Supplied planning examples', 'count': len(rows),
        'workbookSha256': hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
        'datePolicy': 'Original supplied values, normalized for parsing; research annotations do not overwrite them.',
        'geometryPolicy': 'Arithmetic midpoint of available named endpoints, or single endpoint fallback. Approximate representative points, not construction routes.',
        'sourceWorkbook': WORKBOOK_URL,
    }
    return rows, metadata


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook', type=Path, default=ROOT / 'public/sources/Projects_Overlaps.xlsx')
    parser.add_argument('--review', type=Path, default=ROOT / 'data/review_annotations.json')
    parser.add_argument('--check', action='store_true', help='Compare generated content without modifying files.')
    args = parser.parse_args()
    review = json.loads(args.review.read_text(encoding='utf-8'))
    rows, metadata = build_dataset(args.workbook, review)
    outputs = {
        ROOT / 'src/data/projects.json': json.dumps(rows, indent=2, ensure_ascii=False, allow_nan=False) + '\n',
        ROOT / 'src/data/metadata.json': json.dumps(metadata, indent=2, allow_nan=False) + '\n',
    }
    stale = []
    for path, content in outputs.items():
        if args.check:
            if not path.is_file() or path.read_text(encoding='utf-8') != content:
                stale.append(str(path.relative_to(ROOT)))
        else:
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(content, encoding='utf-8')
    if stale:
        parser.exit(1, 'Generated files need refreshing: ' + ', '.join(stale) + '\nRun python scripts/import_workbook.py to refresh them.\n')
    action = 'Verified' if args.check else 'Imported'
    print(f'{action} {len(rows)} supplied projects; originals retained.')


if __name__ == '__main__':
    main()
